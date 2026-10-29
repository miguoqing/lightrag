#!/usr/bin/env python3
"""把单个文档写入本地 LightRAG 知识库。

走 HTTP API（而不是在进程里 new 一个 LightRAG）：本机默认的 NetworkX 图存储是单写者，
服务进程已经把图缓存在内存里并整份重写文件，第二个进程同时写会互相覆盖。交给正在跑的
服务处理，既避开冲突，也复用了服务端配置好的解析引擎和 LLM/embedding。

用法：
    python ingest_file.py inspect <file> [--json] [--base-url URL] [--project-dir DIR]
    python ingest_file.py ingest  <file> [--json] [--no-wait] [--timeout 1800]
                                        [--replace] [--purge-on-failure] [--project-dir DIR]
    python ingest_file.py purge   <file> [--json] [--project-dir DIR]

边界：只读用户指定的文件；写入/删除只发生在项目目录下的 inputs/。
repo-src / lightrag / .venv / rag_storage 一律不碰（见 PROTECTED_DIR_NAMES）。

退出码：0 成功 / 1 入库失败或服务不可用 / 2 输入无效
"""

from __future__ import annotations

import argparse
import http.client
import json
import mimetypes
import os
import re
import shutil
import sys
import time
import uuid
from html.parser import HTMLParser
from pathlib import Path
from urllib import error as urlerror
from urllib import request as urlrequest
from urllib.parse import urlparse

DEFAULT_BASE_URL = os.environ.get("LIGHTRAG_BASE_URL", "http://127.0.0.1:9621")
TEXT_EXTS = {
    ".md", ".mdx", ".txt", ".json", ".yaml", ".yml", ".html", ".htm",
    ".csv", ".xml", ".log", ".ini", ".conf", ".properties",
}
DOC_EXTS = {".pdf", ".docx"}
TERMINAL_STATUSES = {"processed", "failed"}
PROTECTED_DIR_NAMES = {
    "repo-src", "lightrag", ".venv", "rag_storage", ".git", ".skill-build"
}
MAX_ESTIMATE_BYTES = 5 * 1024 * 1024
POLL_INTERVAL = 3.0
PROGRESS_EVERY = 15.0


class UsageError(Exception):
    """输入本身有问题（文件不存在、空文件、后缀不支持）。"""


class ApiError(Exception):
    def __init__(self, status: int | None, message: str, detail: str = ""):
        super().__init__(message)
        self.status = status
        self.message = message
        self.detail = detail


# --------------------------------------------------------------------------- HTTP


def _base(base_url: str | None) -> str:
    return (base_url or DEFAULT_BASE_URL).rstrip("/")


def _request(
    base_url: str,
    path: str,
    *,
    method: str = "GET",
    body: bytes | None = None,
    headers: dict | None = None,
    timeout: float = 60,
) -> dict:
    url = f"{_base(base_url)}/{path.lstrip('/')}"
    hdrs = {"Accept": "application/json"}
    api_key = os.environ.get("LIGHTRAG_API_KEY")
    if api_key:
        hdrs["X-API-Key"] = api_key
    if headers:
        hdrs.update(headers)

    req = urlrequest.Request(url, data=body, headers=hdrs, method=method)
    try:
        with urlrequest.urlopen(req, timeout=timeout) as resp:
            raw = resp.read()
            if not raw:
                return {}
            try:
                return json.loads(raw)
            except json.JSONDecodeError:
                # 代理或网关返回 HTML（登录页、错误页）时，别把 traceback 甩给用户
                raise ApiError(
                    resp.status,
                    f"{url} 返回的不是 JSON",
                    raw[:200].decode("utf-8", "replace"),
                ) from None
    except urlerror.HTTPError as exc:
        raw = exc.read().decode("utf-8", "replace")
        try:
            parsed = json.loads(raw)
            detail = parsed.get("detail") or parsed
            detail = detail if isinstance(detail, str) else json.dumps(detail, ensure_ascii=False)
        except json.JSONDecodeError:
            detail = raw.strip()[:500]
        raise ApiError(exc.code, f"HTTP {exc.code}", detail) from None
    except urlerror.URLError as exc:
        raise ApiError(None, f"无法连接 {_base(base_url)}：{exc.reason}") from None
    except TimeoutError:
        raise ApiError(None, f"请求 {url} 超时") from None


def get_health(base_url: str) -> dict:
    return _request(base_url, "/health", timeout=10)


def get_supported_extensions(base_url: str) -> set[str]:
    data = _request(base_url, "/documents/supported_file_types", timeout=15)
    return {ext.lower() for ext in data.get("supported_extensions", [])}


def get_graph_labels(base_url: str) -> set[str] | None:
    """图谱里现有的全部实体名，用于入库后的“有没有真的建出东西”自检。"""
    try:
        data = _request(base_url, "/graph/label/list", timeout=30)
    except ApiError:
        return None
    if isinstance(data, list):
        return {str(item) for item in data}
    return None


def get_paginated(base_url: str, page: int, page_size: int = 200) -> dict:
    payload = json.dumps({"page": page, "page_size": page_size}).encode("utf-8")
    return _request(
        base_url,
        "/documents/paginated",
        method="POST",
        body=payload,
        headers={"Content-Type": "application/json"},
        timeout=30,
    )


def find_documents_by_name(base_url: str, doc_name: str) -> list[dict]:
    """按 file_path 反查文档记录。

    真实 doc_id 是 `doc-<hash>`，文件名只是重名冲突的判定键——删除前必须先反查。
    """
    matches: list[dict] = []
    page = 1
    while True:
        data = get_paginated(base_url, page)
        documents = data.get("documents") or []
        for doc in documents:
            file_path = str(doc.get("file_path") or "")
            if Path(file_path).name == doc_name:
                matches.append(doc)
        pagination = data.get("pagination") or {}
        total_pages = int(pagination.get("total_pages") or 1)
        if not documents or page >= total_pages:
            return matches
        page += 1


def upload_file(base_url: str, path: Path, timeout: float = 600) -> dict:
    """流式上传：整个文件不进内存（服务端上限 100MB，正好是会把内存打满的量级）。"""
    boundary = "----LightRAGBoundary" + uuid.uuid4().hex
    content_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    head = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="file"; filename="{path.name}"\r\n'
        f"Content-Type: {content_type}\r\n\r\n"
    ).encode("utf-8")
    tail = f"\r\n--{boundary}--\r\n".encode("ascii")
    total = len(head) + path.stat().st_size + len(tail)

    parsed = urlparse(_base(base_url))
    connection_class = (
        http.client.HTTPSConnection if parsed.scheme == "https" else http.client.HTTPConnection
    )
    target = (parsed.path.rstrip("/") + "/documents/upload") or "/documents/upload"
    headers = {
        "Content-Type": f"multipart/form-data; boundary={boundary}",
        "Content-Length": str(total),
        "Accept": "application/json",
    }
    api_key = os.environ.get("LIGHTRAG_API_KEY")
    if api_key:
        headers["X-API-Key"] = api_key

    connection = connection_class(parsed.hostname, parsed.port, timeout=timeout)
    try:
        connection.putrequest("POST", target)
        for name, value in headers.items():
            connection.putheader(name, value)
        connection.endheaders()
        connection.send(head)
        with path.open("rb") as handle:
            while True:
                chunk = handle.read(1024 * 1024)
                if not chunk:
                    break
                connection.send(chunk)
        connection.send(tail)
        response = connection.getresponse()
        raw = response.read()
        status = response.status
    except (OSError, http.client.HTTPException) as exc:
        raise ApiError(None, f"上传失败：{exc}") from None
    finally:
        connection.close()

    if status >= 400:
        raise ApiError(status, f"HTTP {status}", _detail_from(raw))
    if not raw:
        return {}
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        raise ApiError(status, "上传接口返回的不是 JSON", raw[:200].decode("utf-8", "replace")) from None


def _detail_from(raw: bytes) -> str:
    text = raw.decode("utf-8", "replace")
    try:
        parsed = json.loads(text)
        detail = parsed.get("detail") or parsed
        return detail if isinstance(detail, str) else json.dumps(detail, ensure_ascii=False)
    except json.JSONDecodeError:
        return text.strip()[:500]


def get_track_status(base_url: str, track_id: str) -> dict:
    return _request(base_url, f"/documents/track_status/{track_id}", timeout=30)


def delete_document(base_url: str, doc_id: str) -> dict:
    payload = json.dumps(
        {"doc_ids": [doc_id], "delete_file": True, "delete_llm_cache": True}
    ).encode("utf-8")
    return _request(
        base_url,
        "/documents/delete_document",
        method="DELETE",
        body=payload,
        headers={"Content-Type": "application/json"},
        timeout=120,
    )


def _document_exists(base_url: str, doc_id: str) -> bool:
    """全量翻页确认，不能只看第一页——文档多了以后会误判“已删除”。"""
    page = 1
    while True:
        data = get_paginated(base_url, page)
        documents = data.get("documents") or []
        if any(doc.get("id") == doc_id for doc in documents):
            return True
        total_pages = int((data.get("pagination") or {}).get("total_pages") or 1)
        if not documents or page >= total_pages:
            return False
        page += 1


def wait_until_absent(base_url: str, doc_id: str, timeout: float = 90) -> bool:
    deadline = time.time() + timeout
    while time.time() < deadline:
        if not _document_exists(base_url, doc_id):
            return True
        time.sleep(2)
    return False


# ------------------------------------------------------------------ 操作边界


def _is_within(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


def _is_protected(target: Path, project_dir: Path) -> bool:
    """项目目录之外、或落在 repo-src/lightrag/.venv/rag_storage 里的路径，一律不动。"""
    if not _is_within(target, project_dir):
        return True
    relative = target.resolve().relative_to(project_dir.resolve())
    return any(part in PROTECTED_DIR_NAMES for part in relative.parts)


def _resolve_project_dir(explicit: str | None, input_directory: str | None) -> Path:
    candidates = [explicit, os.environ.get("LIGHTRAG_PROJECT_DIR")]
    if input_directory:
        candidates.append(str(Path(input_directory).parent))
    candidates.append(os.getcwd())
    for candidate in candidates:
        if candidate:
            return Path(candidate).resolve()
    return Path.cwd().resolve()


def _backup_source(input_directory: str, path: Path) -> str | None:
    """替换前留一份旧上传副本（仅当服务端还留着这张源文件）。"""
    if not input_directory:
        return None
    source = Path(input_directory) / path.name
    if not source.is_file():
        return None
    try:
        backup_dir = Path(input_directory) / ".replace-backup"
        backup_dir.mkdir(parents=True, exist_ok=True)
        target = backup_dir / f"{path.name}.{int(time.time())}"
        shutil.copy2(source, target)
        return str(target)
    except OSError:
        return None


def _clean_local_artifacts(
    input_directory: str | None, doc_name: str, project_dir: Path
) -> tuple[list[str], list[str]]:
    """删掉解析产物与上传副本；越界或受保护的路径只记录、不删除。"""
    removed: list[str] = []
    skipped: list[str] = []
    if not input_directory:
        return removed, skipped
    root = Path(input_directory)
    if not root.is_dir():
        return removed, skipped

    candidates = [root / doc_name]
    parsed_dir = root / "__parsed__"
    if parsed_dir.is_dir():
        candidates.extend(parsed_dir.glob(f"{Path(doc_name).stem}*"))

    for item in candidates:
        if not item.is_file():
            continue
        if _is_protected(item, project_dir):
            skipped.append(str(item))
            continue
        try:
            item.unlink()
            removed.append(str(item))
        except OSError:
            skipped.append(str(item))
    return removed, skipped


def purge_document(
    base_url: str,
    doc_name: str,
    input_directory: str | None = None,
    *,
    project_dir: str | Path | None = None,
    wait_timeout: float = 90,
) -> dict:
    """把一篇文档在知识库里留下的痕迹清干净。

    顺序：反查真实 doc_id → 删除（连同上传文件与 LLM 缓存）→ 轮询确认记录消失
    → 清理本机解析产物。清理逻辑本身不判断“该不该清”——那是调用方/用户的决定。
    """
    root = _resolve_project_dir(
        str(project_dir) if project_dir is not None else None, input_directory
    )
    report: dict = {
        "doc_name": doc_name,
        "doc_ids": [],
        "record_removed": None,
        "artifacts_removed": [],
        "artifacts_skipped": [],
        "project_dir": str(root),
        "error": None,
    }
    try:
        found = find_documents_by_name(base_url, doc_name)
        report["doc_ids"] = [doc.get("id") for doc in found]
        if not found:
            report["record_removed"] = True  # 本来就没有，等于已清理
        else:
            for doc_id in report["doc_ids"]:
                delete_document(base_url, doc_id)
            report["record_removed"] = all(
                wait_until_absent(base_url, doc_id, wait_timeout)
                for doc_id in report["doc_ids"]
            )
        removed, skipped = _clean_local_artifacts(input_directory, doc_name, root)
        report["artifacts_removed"] = removed
        report["artifacts_skipped"] = skipped
    except ApiError as exc:
        report["error"] = f"{exc.message}：{exc.detail}" if exc.detail else exc.message
    except Exception as exc:  # 清理失败不能掩盖原始错误
        report["error"] = f"{type(exc).__name__}: {exc}"
    return report


# ----------------------------------------------------------------- text / tokens


class _HTMLText(HTMLParser):
    SKIP = {"script", "style", "noscript", "template", "svg"}
    BLOCK = {
        "p", "div", "br", "li", "tr", "td", "th", "h1", "h2", "h3", "h4", "h5",
        "h6", "section", "article", "header", "footer", "pre", "blockquote", "table",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._chunks: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip_depth += 1
        elif tag in self.BLOCK:
            self._chunks.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip_depth:
            self._skip_depth -= 1
        elif tag in self.BLOCK:
            self._chunks.append("\n")

    def handle_data(self, data):
        if not self._skip_depth:
            self._chunks.append(data)

    def text(self) -> str:
        joined = "".join(self._chunks)
        return re.sub(r"[ \t]+\n", "\n", re.sub(r"\n{3,}", "\n\n", joined)).strip()


def _looks_binary(path: Path, sample: int = 8192) -> bool:
    """二进制文件不能当文本读——否则随机字节会被算成几万 token 的“正文”。"""
    try:
        head = path.open("rb").read(sample)
    except OSError:
        return True
    if b"\x00" in head:
        return True
    text = head.decode("utf-8", errors="replace")
    if not text:
        return False
    unprintable = sum(1 for ch in text if not (ch.isprintable() or ch in "\n\r\t"))
    return unprintable / len(text) > 0.1


def extract_text(path: Path) -> tuple[str | None, str, dict]:
    """返回 (文本, 抽取方式, 元信息)；拿不到文本时文本为 None。"""
    ext = path.suffix.lower()
    if ext in {".html", ".htm"}:
        parser = _HTMLText()
        parser.feed(path.read_text(encoding="utf-8", errors="replace"))
        return parser.text(), "html 去标签", {}
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            return None, "pypdf 未安装", {}
        reader = PdfReader(str(path))
        pages = len(reader.pages)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        meta = {"pages": pages}
        if pages and len(text.strip()) < 30 * pages:
            meta["warning"] = (
                f"{pages} 页 PDF 只提取到 {len(text.strip())} 个字符，疑似扫描件（没有文本层）。"
                "入库后内容可能为空或直接失败，建议先 OCR，或换用 docling/mineru 解析引擎"
            )
        return text, "pypdf", meta
    if ext == ".docx":
        try:
            import docx
        except ImportError:
            return None, "python-docx 未安装", {}
        document = docx.Document(str(path))
        text = "\n".join(p.text for p in document.paragraphs)
        meta = {"paragraphs": len(document.paragraphs)}
        if not text.strip():
            meta["warning"] = "docx 里没有段落文本（可能全是图片/表格），入库后内容可能为空"
        return text, "python-docx", meta

    if _looks_binary(path):
        return None, "二进制文件，交给服务端解析", {}

    try:
        return path.read_text(encoding="utf-8", errors="replace"), "按 UTF-8 文本读取", {}
    except OSError as exc:
        return None, f"读取失败：{exc}", {}


def estimate_tokens(text: str) -> tuple[int, str]:
    try:
        import tiktoken

        return len(tiktoken.get_encoding("cl100k_base").encode(text)), "tiktoken / cl100k_base"
    except Exception:
        return max(1, round(len(text) / 3)), "字符数/3（未安装 tiktoken，粗估）"


# ----------------------------------------------------------------------- inspect


def doc_type_for(ext: str) -> str:
    if ext in {".md", ".mdx", ".txt", ".json", ".yaml", ".yml", ".html", ".htm"}:
        return "text"
    if ext in DOC_EXTS:
        return "document"
    return "auto"


def inspect_file(path: Path, base_url: str, project_dir: str | None = None) -> dict:
    if not path.exists():
        raise UsageError(f"文件不存在：{path}")
    if not path.is_file():
        raise UsageError(f"不是普通文件：{path}")
    size = path.stat().st_size
    if size == 0:
        raise UsageError(f"文件是空的：{path}")

    ext = path.suffix.lower()
    if not ext:
        raise UsageError("文件没有扩展名，无法判断类型；请补上扩展名或转成 .md/.txt")

    working_dir = os.environ.get("WORKING_DIR") or ""
    input_dir = ""
    service_ok, service_note = True, ""
    supported: set[str] = set()
    try:
        health = get_health(base_url)
        working_dir = health.get("working_directory") or working_dir
        input_dir = health.get("input_directory") or ""
        supported = get_supported_extensions(base_url)
    except ApiError as exc:
        service_ok = False
        service_note = str(exc)

    if service_ok and supported and ext not in supported:
        raise UsageError(
            f"服务端不接受 {ext}；可用后缀：{', '.join(sorted(supported))}"
        )

    resolved_project_dir = _resolve_project_dir(project_dir, input_dir)

    existing: list[dict] = []
    if service_ok:
        try:
            existing = find_documents_by_name(base_url, path.name)
        except ApiError:
            existing = []

    if size > MAX_ESTIMATE_BYTES:
        text, method, meta = None, "文件过大，未做本地解析", {}
    else:
        text, method, meta = extract_text(path)

    usable = bool(text and text.strip())
    if usable:
        tokens, token_method = estimate_tokens(text)
    else:
        tokens, token_method = None, "无法本地预估：没有可读文本"

    warnings: list[str] = []
    if meta.get("warning"):
        warnings.append(meta["warning"])
    if not usable:
        warnings.append(
            f"没有提取到可读文本（{method}），无法预估 token；"
            "实际内容由服务端解析器决定，入库后可能为空或失败"
        )
    if existing:
        warnings.append(
            f"同名文档已存在（{existing[0].get('id')}，{existing[0].get('status')}），"
            "上传会 409；用 --replace 替换，或改名后再入库"
        )
    warning = "；".join(warnings) or None

    return {
        "file": str(path.resolve()),
        "doc_name": path.name,
        "doc_type": doc_type_for(ext),
        "extension": ext,
        "source": "本地文件",
        "size_bytes": size,
        "size_human": _human_size(size),
        "content_chars": len(text) if text else None,
        "estimated_tokens": tokens,
        "token_method": token_method,
        "text_method": method,
        "text_meta": meta,
        "warning": warning,
        "has_conflict": bool(existing),
        "existing": [
            {"id": doc.get("id"), "status": doc.get("status")} for doc in existing
        ],
        "working_dir": working_dir or "（未知：服务未连接且未设置 WORKING_DIR）",
        "input_dir": input_dir,
        "project_dir": str(resolved_project_dir),
        "base_url": _base(base_url),
        "service_ok": service_ok,
        "service_note": service_note,
    }


def _human_size(num: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if num < 1024 or unit == "GB":
            return f"{num:.1f} {unit}" if unit != "B" else f"{num} B"
        num /= 1024
    return f"{num:.1f} GB"


def print_inspect_human(info: dict) -> None:
    chars = (
        f"，正文 {info['content_chars']:,} 字符"
        if (info["content_chars"] or 0) >= 20
        else ""
    )
    tokens = (
        f"≈ {info['estimated_tokens']:,}" if info["estimated_tokens"] is not None else "未知"
    )
    print(f"文件：{info['doc_name']}")
    print(f"文档类型：{info['doc_type']}")
    print(f"来源：{info['file']}（{info['source']}）")
    print(f"大小：{info['size_human']}")
    print(f"预估 token：{tokens}（{info['token_method']}{chars}）")
    print(f"目标知识库：{info['working_dir']}")
    if info.get("warning"):
        print(f"⚠ 提示：{info['warning']}")
    if not info["service_ok"]:
        print(f"⚠ 服务未连接：{info['service_note']}")


# ------------------------------------------------------------------------ ingest


def _status_of(doc: dict) -> str:
    return str(doc.get("status") or "").lower()


def wait_for_document(
    base_url: str, track_id: str, timeout: float, *, quiet: bool = False
) -> dict | None:
    deadline = time.time() + timeout
    started = time.time()
    last_progress = started
    last: dict | None = None
    while time.time() < deadline:
        data = get_track_status(base_url, track_id)
        documents = data.get("documents") or []
        if documents:
            last = documents[0]
            if _status_of(last) in TERMINAL_STATUSES:
                return last
        now = time.time()
        if not quiet and now - last_progress >= PROGRESS_EVERY:
            elapsed = int(now - started)
            status = _status_of(last) if last else "等待入队"
            print(f"  处理中… 状态={status or '未知'}，已用 {elapsed}s", file=sys.stderr, flush=True)
            last_progress = now
        time.sleep(POLL_INTERVAL)
    return last


def upload_file_with_retry(base_url: str, path: Path, attempts: int = 3) -> dict:
    """删除是异步的：刚删完立刻上传偶尔还会撞 409，退避重试几次。"""
    last_error: ApiError | None = None
    for attempt in range(attempts):
        try:
            return upload_file(base_url, path)
        except ApiError as exc:
            if exc.status != 409 or attempt == attempts - 1:
                raise
            last_error = exc
            time.sleep(1.5 * (attempt + 1))
    raise last_error or ApiError(None, "上传失败")


def ingest_file(
    path: Path,
    base_url: str,
    *,
    wait: bool,
    timeout: float,
    replace: bool,
    purge_on_failure: bool = False,
    quiet: bool = False,
    project_dir: str | None = None,
) -> dict:
    info = inspect_file(path, base_url, project_dir)
    if not info["service_ok"]:
        raise ApiError(None, f"服务不可用：{info['service_note']}")

    labels_before = get_graph_labels(base_url) if wait else None
    purge_report: dict | None = None
    backup_path: str | None = None
    try:
        response = upload_file(base_url, path)
    except ApiError as exc:
        if exc.status != 409 or not replace:
            raise
        # 替换是“先删后传”：先留一份旧副本，删完传失败时至少还有东西可恢复
        backup_path = _backup_source(info.get("input_dir") or "", path)
        purge_report = purge_document(
            base_url, path.name, info.get("input_dir"),
            project_dir=info.get("project_dir"),
        )
        if purge_report.get("record_removed") is False:
            raise ApiError(
                None, f"同名文档删除失败，无法替换：{purge_report.get('error')}"
            )
        labels_before = get_graph_labels(base_url) if wait else None
        try:
            response = upload_file_with_retry(base_url, path)
        except ApiError as replacement_error:
            hint = (
                f"旧副本已备份到 {backup_path}"
                if backup_path
                else "旧文档已删除且没有可用的本地副本"
            )
            raise ApiError(
                replacement_error.status, f"替换失败：{hint}", replacement_error.detail
            ) from None

    track_id = response.get("track_id") or ""
    result = {
        **info,
        "stage": "upload",
        "track_id": track_id,
        "doc_id": None,
        "status": "queued",
        "chunks_count": None,
        "error": None,
        "message": response.get("message", ""),
        "purge": purge_report,
        "backup": backup_path,
        "graph": None,
    }
    if not wait or not track_id:
        return result

    result["stage"] = "wait"
    document = wait_for_document(base_url, track_id, timeout, quiet=quiet)
    if document is None:
        result.update(status="timeout", error="等待超时，文档可能仍在处理")
    else:
        result.update(
            status=_status_of(document),
            doc_id=document.get("id"),
            chunks_count=document.get("chunks_count"),
            error=document.get("error_msg") or document.get("error"),
            file_path=document.get("file_path"),
        )

    if result["status"] == "processed":
        result["stage"] = "done"
        labels_after = get_graph_labels(base_url)
        if labels_before is not None and labels_after is not None:
            result["graph"] = {
                "entities_before": len(labels_before),
                "entities_after": len(labels_after),
                "entities_added": len(labels_after - labels_before),
            }
    elif purge_on_failure:
        result["purge"] = purge_document(
            base_url, path.name, info.get("input_dir"),
            project_dir=info.get("project_dir"),
        )
    return result


def print_ingest_human(result: dict) -> None:
    if result["status"] == "queued":
        print(f"已提交：{result['doc_name']}")
        print(f"文档类型：{result['doc_type']}")
        print(f"存储路径：{result['working_dir']}")
        print(f"track_id：{result['track_id']}")
        print("状态：服务端正在后台处理，尚未完成")
        print(f"下一步：用 /documents/track_status/{result['track_id']} 查看结果，或在 WebUI 文档列表看状态。")
        return

    if result["status"] == "timeout":
        print(f"处理未完成：{result['doc_name']}")
        print(f"失败步骤：{result.get('stage') or '未知'}")
        print(f"状态：等待超时（{result.get('error') or '仍在处理'}）")
        print("建议：稍后在 WebUI 文档列表确认最终状态；需要继续等待时加 --timeout 3600 重新提交（同名会被拒绝，先删旧文档）")
        _print_purge(result)
        return

    if result["status"] == "processed":
        print(f"入库完成：{result['doc_name']}")
        print(f"文档类型：{result['doc_type']}")
        print(f"存储路径：{result['working_dir']}")
        print(f"文档 ID：{result['doc_id']}")
        if result.get("chunks_count") is not None:
            print(f"分块数：{result['chunks_count']}")
        graph = result.get("graph")
        if graph:
            print(
                f"图谱：新增实体 {graph['entities_added']}"
                f"（{graph['entities_before']} → {graph['entities_after']}）"
            )
            if graph["entities_added"] == 0:
                print("⚠ 图谱没有新增实体：可能抽取为空、内容重复，或实体已全部存在")
        if result.get("backup"):
            print(f"旧副本：{result['backup']}")
        print("下一步：可以检索知识，或生成测试用例。")
        return

    print(f"入库失败：{result['doc_name']}")
    print(f"失败步骤：{result.get('stage') or '未知'}")
    print(f"状态：{result['status']}")
    print(f"原因：{result.get('error') or result.get('message') or '未知'}")
    print(f"建议：{_suggest(result)}")
    if result.get("backup"):
        print(f"旧副本：{result['backup']}")
    _print_retry(result)
    _print_purge(result)


def _print_retry(result: dict) -> None:
    file = result.get("file") or result["doc_name"]
    print("重试命令：")
    print(f'  python ingest_file.py ingest "{file}" --timeout 3600')
    print(f'  python ingest_file.py ingest "{file}" --replace --timeout 3600   # 同名冲突时')


def _print_purge(result: dict) -> None:
    purge = result.get("purge")
    if purge is None:
        print("清理：尚未执行 —— 需要你确认")
        print(f'  清理这篇文档：python ingest_file.py purge "{result.get("file", result["doc_name"])}"')
        print("  保留现场：不做任何操作即可；之后也可以随时再清理")
        return
    if purge.get("error"):
        print(f"清理：失败 —— {purge['error']}（请手动检查 doc_status 与 inputs/__parsed__）")
        return
    ids = "、".join(purge.get("doc_ids") or []) or "无"
    record = "已删除" if purge.get("record_removed") else "⚠ 仍存在"
    artifacts = purge.get("artifacts_removed") or []
    skipped = purge.get("artifacts_skipped") or []
    line = f"清理：文档记录 {record}（{ids}），解析残留清除 {len(artifacts)} 个"
    if skipped:
        line += f"，另有 {len(skipped)} 个因不在允许目录内被跳过"
    print(line)


def _suggest(result: dict) -> str:
    error = str(result.get("error") or "")
    lowered = error.lower()
    if "already exists" in lowered or "duplicate" in lowered:
        return "同名文档已存在：用 --replace 自动替换，或先手动删除旧文档"
    if "content" in lowered and "exists" in lowered:
        return "内容重复：知识库里已有相同内容，确认是否还需要这份副本"
    if "timeout" in lowered or "超时" in error:
        return "处理超时：用 --timeout 3600 重试，或稍后在 WebUI 看这篇文档的最终状态"
    return "查看服务端日志（默认 lightrag.log，本机可能已挪到 repo-src/lightrag.log）中同时间的堆栈；模型/embedding 配置或文件损坏是最常见原因"


# --------------------------------------------------------------------------- CLI


def _add_connection_options(parser: argparse.ArgumentParser, *, in_subcommand: bool = False) -> None:
    """--base-url / --project-dir 放在子命令前后都能用。

    argparse 的顶层选项默认只能写在子命令前面；但人（和模型）习惯写成
    `inspect file --base-url ...`。子命令里用 SUPPRESS 兜一份，避免把顶层的值覆盖成 None。
    """
    default = argparse.SUPPRESS if in_subcommand else None
    parser.add_argument(
        "--base-url", default=default, help=f"LightRAG 地址（默认 {DEFAULT_BASE_URL}）"
    )
    parser.add_argument(
        "--project-dir",
        default=default,
        help="项目目录（默认 $LIGHTRAG_PROJECT_DIR，否则取 inputs 的父目录或当前目录）",
    )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="把文档写入本地 LightRAG 知识库")
    _add_connection_options(parser)
    sub = parser.add_subparsers(dest="command", required=True)

    inspect_cmd = sub.add_parser("inspect", help="只看不写：类型、大小、预估 token、目标路径")
    inspect_cmd.add_argument("file")
    inspect_cmd.add_argument("--json", action="store_true", help="输出 JSON")
    _add_connection_options(inspect_cmd, in_subcommand=True)

    ingest_cmd = sub.add_parser("ingest", help="上传并等待处理完成；失败默认不清理（先上报）")
    ingest_cmd.add_argument("file")
    ingest_cmd.add_argument("--json", action="store_true", help="输出 JSON")
    ingest_cmd.add_argument("--no-wait", action="store_true", help="只上传，不等待处理结果")
    ingest_cmd.add_argument("--timeout", type=float, default=1800, help="等待上限（秒），默认 1800")
    ingest_cmd.add_argument("--replace", action="store_true", help="同名冲突时先删除旧文档再上传")
    ingest_cmd.add_argument(
        "--purge-on-failure",
        action="store_true",
        help="失败时自动清理（默认不清理：先上报，由用户决定；仅自动化流程使用）",
    )
    _add_connection_options(ingest_cmd, in_subcommand=True)

    purge_cmd = sub.add_parser("purge", help="清理一篇文档在知识库里的全部痕迹")
    purge_cmd.add_argument("file")
    purge_cmd.add_argument("--json", action="store_true", help="输出 JSON")
    _add_connection_options(purge_cmd, in_subcommand=True)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    path = Path(args.file).expanduser()

    try:
        if args.command == "inspect":
            info = inspect_file(path, args.base_url, args.project_dir)
            if args.json:
                print(json.dumps(info, ensure_ascii=False, indent=2))
            else:
                print_inspect_human(info)
            return 0

        if args.command == "purge":
            # 清理不要求本地文件还在：失败后用户可能已经删掉了源文件
            try:
                input_dir = get_health(args.base_url).get("input_directory") or ""
            except ApiError:
                input_dir = ""
            report = purge_document(
                args.base_url, path.name, input_dir, project_dir=args.project_dir
            )
            if args.json:
                print(json.dumps(report, ensure_ascii=False, indent=2))
            else:
                ids = "、".join(report.get("doc_ids") or []) or "未找到"
                print(f"清理对象：{path.name}")
                record = "已删除" if report.get("record_removed") else "⚠ 仍存在"
                print(f"文档记录：{record}（{ids}）")
                print(f"解析残留：清除 {len(report.get('artifacts_removed') or [])} 个")
                skipped = report.get("artifacts_skipped") or []
                if skipped:
                    print(f"跳过：{len(skipped)} 个（不在允许目录内）")
                if report.get("error"):
                    print(f"⚠ 清理出错：{report['error']}")
            return 0 if not report.get("error") else 1

        result = ingest_file(
            path,
            args.base_url,
            wait=not args.no_wait,
            timeout=args.timeout,
            replace=args.replace,
            purge_on_failure=args.purge_on_failure,
            quiet=args.json,
            project_dir=args.project_dir,
        )
        if args.json:
            print(json.dumps(result, ensure_ascii=False, indent=2))
        else:
            print_ingest_human(result)
        return 0 if result["status"] in {"processed", "queued"} else 1

    except UsageError as exc:
        print(f"输入无效：{exc}", file=sys.stderr)
        return 2
    except ApiError as exc:
        message = exc.message + (f"：{exc.detail}" if exc.detail else "")
        print(f"请求失败：{message}", file=sys.stderr)
        if exc.status == 409:
            print("提示：同名文档已存在，用 --replace 替换，或先删除旧文档", file=sys.stderr)
        elif exc.status in (401, 403):
            print("提示：服务开启了鉴权，请导出 LIGHTRAG_API_KEY", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
